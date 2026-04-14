"""
Apollo Contact Fetcher — runs directly with company list from Apollo search.
No input Excel needed: company names are embedded from the company-search step.
Writes apollo_contacts_output.xlsx in the current directory.
"""

import time
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

API_KEY   = "AQ-JughbK6q6NAGChDI4OA"
OUT_FILE  = "apollo_contacts_output.xlsx"
DELAY     = 0.6   # seconds between calls

SEARCH_URL = "https://api.apollo.io/api/v1/mixed_people/search"
ENRICH_URL = "https://api.apollo.io/api/v1/people/match"

SENIORITIES = ["c_suite", "vp", "director", "manager"]
TITLES = [
    "CFO", "COO", "CTO", "CEO",
    "VP Finance", "VP Sales", "VP Operations",
    "Director Finance", "Director Sales", "Director Operations",
    "Sales Manager", "Finance Manager", "Operations Manager",
    "Head of Finance", "Head of Sales", "Head of Operations",
]

COLUMNS = [
    "Company", "First Name", "Last Name", "Title",
    "Seniority", "Department", "Email", "LinkedIn URL", "Phone",
]

# 25 Swedish geothermal equipment companies fetched from Apollo company search
COMPANIES = [
    "Thermia AB",
    "Qvantum Sverige",
    "MuoviTech",
    "Accio AB",
    "DREM",
    "Aira",
    "Baseload Capital",
    "Climeon",
    "NIBE",
    "Malmberg",
    "Samster AB",
    "Azelio",
    "Kustens vvs AB",
    "Borrspecialisten AB",
    "VÄRNAMO - FORSHEDA RÖR AB",
    "Absolicon Solar Collector AB",
    "Hydroc Energy",
    "Akademiska Hus",
    "Againity",
    "Systemair Group",
    "Swedish Refrigeration & Heat Pump Association",
    "International Ground Source Heat Pump Association - Swedish chapter",
    "Gullspång Invest",
    "GRILLBY & F100 RÖR AB",
    "HPT TCP - Technology Collaboration Programme on Heat Pumping Technologies by IEA",
]


def search_people(company: str) -> list | None:
    payload = {
        "api_key": API_KEY,
        "q_keywords": company,
        "person_locations": ["Sweden"],
        "person_seniorities": SENIORITIES,
        "person_titles": TITLES,
        "per_page": 5,
        "page": 1,
    }
    try:
        r = requests.post(SEARCH_URL, json=payload, timeout=30)
        r.raise_for_status()
        return r.json().get("people", [])
    except requests.HTTPError as e:
        print(f"  HTTP {e.response.status_code}")
        return None
    except Exception as e:
        print(f"  Error: {e}")
        return None


def enrich_person(person_id: str) -> dict:
    try:
        r = requests.post(ENRICH_URL, json={"api_key": API_KEY, "id": person_id}, timeout=30)
        r.raise_for_status()
        return r.json().get("person") or {}
    except Exception:
        return {}


def first_phone(p: dict) -> str:
    phones = p.get("phone_numbers") or []
    return phones[0].get("sanitized_number") or phones[0].get("raw_number") or "" if phones else ""


def make_row(company: str, raw: dict, enriched: dict) -> dict:
    s = enriched if enriched else raw
    return {
        "Company":      company,
        "First Name":   s.get("first_name")   or raw.get("first_name")   or "",
        "Last Name":    s.get("last_name")     or raw.get("last_name")    or "",
        "Title":        s.get("title")         or raw.get("title")        or "",
        "Seniority":    s.get("seniority")     or raw.get("seniority")    or "",
        "Department":   s.get("department")    or raw.get("department")   or "",
        "Email":        s.get("email")         or raw.get("email")        or "",
        "LinkedIn URL": s.get("linkedin_url")  or raw.get("linkedin_url") or "",
        "Phone":        first_phone(enriched)  or first_phone(raw),
    }


def write_excel(contacts: list, path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contacts"

    hdr_font  = Font(bold=True, color="FFFFFF")
    hdr_fill  = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    center    = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20

    for ci, col in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = center

    for ri, row in enumerate(contacts, 2):
        for ci, col in enumerate(COLUMNS, 1):
            ws.cell(row=ri, column=ci, value=row.get(col, ""))

    for ci, col in enumerate(COLUMNS, 1):
        cl = get_column_letter(ci)
        ml = max(len(col), *(len(str(ws.cell(r, ci).value or "")) for r in range(2, len(contacts)+2)))
        ws.column_dimensions[cl].width = min(ml + 3, 55)

    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    total   = len(COMPANIES)
    all_contacts = []
    zeros   = []
    errors  = []

    print(f"Processing {total} companies...\n")

    for idx, company in enumerate(COMPANIES, 1):
        print(f"[{idx:2d}/{total}] {company}", end=" → ", flush=True)

        people = search_people(company)
        time.sleep(DELAY)

        if people is None:
            errors.append(company)
            print("ERROR (skipped)")
            continue
        if not people:
            zeros.append(company)
            print("0 contacts")
            continue

        rows = []
        for person in people[:5]:
            pid = person.get("id")
            enriched = {}
            if pid:
                enriched = enrich_person(pid)
                time.sleep(DELAY)
            rows.append(make_row(company, person, enriched))

        all_contacts.extend(rows)
        print(f"{len(rows)} contact(s)")

    write_excel(all_contacts, OUT_FILE)
    print(f"\nSaved → {OUT_FILE}  ({len(all_contacts)} rows)\n")
    print("══ SUMMARY ═══════════════════════════════")
    print(f"  Companies searched : {total}")
    print(f"  Total contacts     : {len(all_contacts)}")
    print(f"  Zero results       : {len(zeros)}")
    for c in zeros:
        print(f"    • {c}")
    print(f"  Errors skipped     : {len(errors)}")
    for c in errors:
        print(f"    • {c}")
    print("══════════════════════════════════════════")


if __name__ == "__main__":
    main()
